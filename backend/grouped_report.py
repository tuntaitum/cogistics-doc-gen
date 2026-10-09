"""
grouped_report.py — one multi-sheet Excel workbook in, one PDF per group out.

First use: Food Pipeline's daily delivery report. The TMS exports one workbook
containing every client's orders, split across sheets by delivery status;
this module produces one report per client.

Driven entirely by a GroupedReportConfig (schemas.py). Nothing here knows
about "Food Pipeline", "hirers" or Thai — those live in the config JSON — so
the same code serves any future "group these rows by a column" report.

Rendering reuses engine.py (fonts, branded page canvas, table builder) so
reports look like the rest of the app.

Design rule: NEVER REPORT FEWER PROBLEMS THAN EXIST. Anything unexpected —
a blank client, an odd status, a different date — is kept in the report and
raised as a warning. A missing sheet or column is a hard error, because
quietly treating it as "no failures today" would tell a client something false.
"""

import os
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from typing import Optional
from xml.sax.saxutils import escape as xml_escape

import openpyxl
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer

import engine
from schemas import DocumentConfig, GroupedReportConfig, ReportSection


class ReportError(Exception):
    """The workbook can't be turned into a trustworthy report (message is user-facing)."""


@dataclass
class ClientReport:
    name: str
    filename: str
    path: str
    counts: dict          # {section key: number of rows}
    total: int


@dataclass
class ReportResult:
    report_date: Optional[date]
    clients: list = field(default_factory=list)     # list[ClientReport]
    warnings: list = field(default_factory=list)    # list[str], shown to the user


# ─────────────────────────────────────────────
#  SMALL HELPERS
# ─────────────────────────────────────────────

def norm(value) -> str:
    """Collapse all whitespace (incl. line breaks) to single spaces. Headers are
    matched this way, so "สินค้าทั้งหมด \\n( รายการ )" equals "สินค้าทั้งหมด ( รายการ )"."""
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def strip_sheet_prefix(name: str) -> str:
    """'2.จัดส่งสำเร็จ' -> 'จัดส่งสำเร็จ' (the TMS numbers its sheets)."""
    return norm(re.sub(r"^\s*\d+\s*\.\s*", "", name))


def clean_cell(value) -> str:
    """Excel cell -> display text. 4.0 -> '4', 15.2 -> '15.2', None -> ''."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return f"{value:.3f}".rstrip("0").rstrip(".")
    return norm(value) if isinstance(value, str) else str(value)


def parse_date(value) -> Optional[date]:
    """Date part of '06-10-2026 (16:00)' (DAY first, as the TMS writes it).
    Also accepts a real Excel date/datetime. Anything else -> None."""
    if hasattr(value, "date") and callable(value.date):      # datetime
        return value.date()
    if isinstance(value, date):
        return value
    m = re.match(r"\s*(\d{1,2})-(\d{1,2})-(\d{4})", str(value or ""))
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def safe_filename(text: str, used: set) -> str:
    """Make `text` safe as a file name on Windows/Mac/Linux and unique within `used`."""
    name = _ILLEGAL_FILENAME_CHARS.sub("-", text)
    name = re.sub(r"\s+", " ", name).strip(" .-") or "report"
    name = name[:150].rstrip(" .-")
    candidate, n = name, 2
    while candidate.casefold() in used:
        candidate = f"{name} ({n})"
        n += 1
    used.add(candidate.casefold())
    return candidate + ".pdf"


def _format_item(item: str, col) -> str:
    """One item of a multi-item cell: keep only the wanted "label value" parts."""
    if not col.part_separator:
        return item
    parts = [p.strip() for p in item.split(col.part_separator) if p.strip()]
    if col.keep_labels is not None:
        wanted = {norm(l) for l in col.keep_labels}
        kept = [p for p in parts if p.split(" ", 1)[0] in wanted]
        if not kept:
            return item                      # unrecognised wording: show everything rather than nothing
        parts = kept
    if not col.show_labels:
        parts = [p.split(" ", 1)[1] if " " in p else p for p in parts]
    return f" {col.part_separator.strip()} ".join(parts)


def _cell_markup(text: str, col) -> str:
    """Text -> safe reportlab Paragraph markup. build_table passes cell strings
    straight into Paragraph, so anything like '&' or '<' must be escaped here,
    and multi-item cells become one item per line."""
    items = [i.strip() for i in text.split(col.split_on) if i.strip()] if col.split_on else [text]
    return "<br/>".join(xml_escape(_format_item(i, col)) for i in items)


# ─────────────────────────────────────────────
#  READING
# ─────────────────────────────────────────────

def _find_sheet(wb, sheet_name: str):
    wanted = norm(sheet_name)
    for ws in wb.worksheets:
        if strip_sheet_prefix(ws.title) == wanted:
            return ws
    raise ReportError(
        f"Sheet '{sheet_name}' was not found in this workbook. "
        f"Sheets present: {', '.join(repr(ws.title) for ws in wb.worksheets)}. "
        "If the TMS renamed a sheet, update the report config."
    )


def _read_section(wb, section: ReportSection, cfg: GroupedReportConfig, warnings: list) -> list:
    """Rows of one sheet as dicts: the section's column keys (display markup)
    plus _group, _date, _status, _row."""
    ws = _find_sheet(wb, section.sheet_name)

    headers = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(cfg.header_row, c).value)
        if h and h not in headers:
            headers[h] = c

    # Headers the report cannot work without -> hard error
    required = [("client", cfg.group_by_header)]
    if cfg.report_date_header:
        required.append(("report date", cfg.report_date_header))
    if section.expected_status:
        required.append(("status", cfg.status_header))
    required += [(f"column '{c.label}'", c.source_header) for c in section.columns if not c.optional]
    missing = [f"{what}: '{norm(h)}'" for what, h in required if norm(h) not in headers]
    if missing:
        raise ReportError(
            f"Sheet '{ws.title}' is missing expected column(s): {'; '.join(missing)}. "
            f"The TMS export format may have changed. (Header row read: row {cfg.header_row}.)"
        )

    col_idx = {}
    for c in section.columns:
        idx = headers.get(norm(c.source_header))
        if idx is None:      # only possible for optional columns
            warnings.append(f"Sheet '{ws.title}': optional column '{c.label}' "
                            f"('{norm(c.source_header)}') not found; it will be left out.")
        col_idx[c.key] = idx

    g_idx = headers[norm(cfg.group_by_header)]
    d_idx = headers.get(norm(cfg.report_date_header)) if cfg.report_date_header else None
    s_idx = headers.get(norm(cfg.status_header)) if cfg.status_header else None

    rows = []
    for r in range(cfg.data_start_row, ws.max_row + 1):
        if all(ws.cell(r, c).value in (None, "") for c in range(1, ws.max_column + 1)):
            continue                                   # blank spacer row
        row = {"_row": r, "_sheet": ws.title}
        row["_group"] = clean_cell(ws.cell(r, g_idx).value)
        row["_date"] = parse_date(ws.cell(r, d_idx).value) if d_idx else None
        row["_status"] = clean_cell(ws.cell(r, s_idx).value) if s_idx else ""
        for c in section.columns:
            idx = col_idx[c.key]
            text = clean_cell(ws.cell(r, idx).value) if idx else ""
            row[c.key] = _cell_markup(text, c)
        rows.append(row)
    return rows


# ─────────────────────────────────────────────
#  RENDERING (one client)
# ─────────────────────────────────────────────

def _render_client(group: str, rows_by_section: dict, cfg: GroupedReportConfig,
                   report_date: Optional[date], out_path: str, assets_dir: str):
    engine._ensure_fonts_registered()
    brand = cfg.brand
    c_primary, c_accent = HexColor(brand.primary), HexColor(brand.accent)
    c_light_bg, c_dark, c_mid = HexColor(brand.light_bg), HexColor(brand.text_dark), HexColor(brand.text_mid)
    c_white = colors.white
    styles = engine.make_styles(c_dark, c_mid, c_white, c_primary)

    total = sum(len(v) for v in rows_by_section.values())
    fmt = dict(
        group=xml_escape(group), total=total,
        date=report_date.strftime("%d-%m-%Y") if report_date else "",
        date_iso=report_date.isoformat() if report_date else "",
    )

    has_signature = any(b.type == "signature" for b in cfg.footer_blocks)
    bottom_margin = (20 + engine.SIGNATURE_AREA_HEIGHT / mm) * mm if has_signature else 20 * mm
    doc = SimpleDocTemplate(
        out_path, pagesize=engine.page_size(cfg),
        leftMargin=engine.MARGIN, rightMargin=engine.MARGIN,
        topMargin=30 * mm, bottomMargin=bottom_margin,
        title=f"{cfg.document_title} - {group}", author=brand.company_name, compress=1,
    )

    summary = "   |   ".join(f"{sec.title}: {len(rows_by_section[sec.key])}" for sec in cfg.sections)
    story = [
        Spacer(1, 4 * mm),
        Paragraph(xml_escape(cfg.document_title), styles["cat_title"]),
        Paragraph(cfg.intro_text_template.format(**fmt), styles["intro"]),
        Paragraph(xml_escape(summary), styles["intro"]),
    ]

    for sec in cfg.sections:
        rows = rows_by_section[sec.key]
        if not rows and not sec.show_when_empty:
            continue
        heading = ParagraphStyle(
            f"Section_{sec.key}", fontName=engine.FONT_BOLD, fontSize=12,
            textColor=HexColor(sec.title_color) if sec.title_color else c_primary,
            spaceBefore=8, spaceAfter=6, leftIndent=-(7 * mm),
        )
        story.append(CondPageBreak(45 * mm))     # heading + at least the first rows must fit, else new page
        story.append(Paragraph(xml_escape(f"{sec.title} ({len(rows)})"), heading))
        if rows:
            # build_table reads .columns / .brand / .orientation from a DocumentConfig;
            # give it one that holds just this section's columns.
            section_cfg = DocumentConfig(
                id=f"{cfg.id}:{sec.key}", name=sec.title, document_title=cfg.document_title,
                columns=list(sec.columns), brand=cfg.brand, orientation=cfg.orientation,
            )
            story.append(engine.build_table(rows, section_cfg, styles, c_primary, c_accent,
                                            c_light_bg, c_mid, c_dark, c_white, auto_height=True))
        else:
            story.append(Paragraph("No orders.", styles["dim_text"]))
        story.append(Spacer(1, 4 * mm))

    story.extend(engine.build_footer_blocks(cfg, styles, c_mid, c_dark))
    doc.build(story, canvasmaker=engine.make_branded_canvas(cfg, assets_dir))


# ─────────────────────────────────────────────
#  PUBLIC ENTRY POINT
# ─────────────────────────────────────────────

def generate_grouped_reports(xlsx_path: str, cfg: GroupedReportConfig,
                             out_dir: str, assets_dir: str) -> ReportResult:
    """Read every section's sheet, group rows by cfg.group_by_header, and write
    one PDF per group into out_dir. Raises ReportError if the workbook can't be
    trusted (missing sheet/column); everything else becomes result.warnings."""
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        raise ReportError(f"Could not open the workbook: {e}")

    warnings: list = []
    try:
        by_section = {sec.key: _read_section(wb, sec, cfg, warnings) for sec in cfg.sections}
    finally:
        wb.close()      # rows are copied out above; don't hold the parsed workbook while rendering PDFs
    all_rows = [r for rows in by_section.values() for r in rows]
    result = ReportResult(report_date=None, warnings=warnings)

    if not all_rows:
        warnings.append("No orders were found on any of the report sheets, so no reports were created.")
        return result

    # ---- report date: the most common date wins (ties -> earliest); odd rows are flagged
    if cfg.report_date_header:
        dated = Counter(r["_date"] for r in all_rows if r["_date"])
        if dated:
            top = max(dated.values())
            result.report_date = min(d for d, n in dated.items() if n == top)
            for d, n in sorted(dated.items()):
                if d != result.report_date:
                    clients = sorted({r["_group"] or cfg.blank_group_label for r in all_rows if r["_date"] == d})
                    warnings.append(
                        f"{n} row(s) have start date {d.strftime('%d-%m-%Y')} instead of the report date "
                        f"{result.report_date.strftime('%d-%m-%Y')} (client(s): {', '.join(clients)}). "
                        "They are included in the reports."
                    )
        undated = sum(1 for r in all_rows if not r["_date"])
        if undated:
            warnings.append(f"{undated} row(s) have a missing or unreadable start date. They are included.")

    # ---- flag (never drop) rows whose status doesn't match what their sheet should hold
    for sec in cfg.sections:
        if sec.expected_status:
            for r in by_section[sec.key]:
                if norm(r["_status"]) != norm(sec.expected_status):
                    warnings.append(
                        f"Sheet '{r['_sheet']}' row {r['_row']}: status is '{r['_status'] or '(blank)'}' "
                        f"but this sheet should only hold '{sec.expected_status}' "
                        f"(client: {r['_group'] or cfg.blank_group_label})."
                    )

    # ---- group. A blank client is kept, under a visible label, and flagged.
    groups: dict = {}
    blank = 0
    for sec in cfg.sections:
        for r in by_section[sec.key]:
            name = r["_group"]
            if not name:
                blank += 1
                name = cfg.blank_group_label
            groups.setdefault(name, {s.key: [] for s in cfg.sections})[sec.key].append(r)
    if blank:
        warnings.append(f"{blank} row(s) have no value in '{norm(cfg.group_by_header)}'. "
                        f"They are reported together under '{cfg.blank_group_label}'; check before sending.")

    # ---- one PDF per group
    os.makedirs(out_dir, exist_ok=True)
    used: set = set()
    fmt_date = dict(date=result.report_date.strftime("%d-%m-%Y") if result.report_date else "",
                    date_iso=result.report_date.isoformat() if result.report_date else "")
    for name in sorted(groups, key=str.casefold):
        rows_by_section = groups[name]
        total = sum(len(v) for v in rows_by_section.values())
        filename = safe_filename(cfg.filename_template.format(group=name, total=total, **fmt_date), used)
        path = os.path.join(out_dir, filename)
        _render_client(name, rows_by_section, cfg, result.report_date, path, assets_dir)
        result.clients.append(ClientReport(
            name=name, filename=filename, path=path, total=total,
            counts={k: len(v) for k, v in rows_by_section.items()},
        ))
    return result
