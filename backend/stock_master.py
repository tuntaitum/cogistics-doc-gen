"""
stock_master.py — product master data for the 3PL stock movement report.

The WMS stock export has no product names, no carton sizes, and no client name. All three come
from here: the Lark Base "3PL Master Data" product list, exported to Excel and uploaded to
CoDocuments. This module:

  * reads and VALIDATES that export (parse_master_workbook),
  * keeps it on the volume, with the previous copy as a safety net (MasterStore),
  * checks that a stock file really belongs to the client someone selected (check_client),
  * turns pieces into "cartons + leftover pieces" (to_cartons / format_cartons / carton_view).

Units, because they are easy to confuse:
  * Everything in the WMS file is PIECES (its "CT" label is wrong).
  * A CARTON is how stock is received and stored: master column "Unit per Storage Unit",
    e.g. "30 pieces per carton". This is what the report uses.
  * A CRATE / basket is how stock is packed for 7-Eleven DCs: "Unit per Delivery Unit". The
    stock report does NOT use it.

IDs are TEXT, matched exactly. "0200639" (a CJ product, from 7-Eleven's CP All order) keeps its
zero. We never turn an ID into a number; if a file's ID has lost its zero we only ADD A HINT.
"""

import json
import math
import os
import re
import shutil
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Optional

import openpyxl

from stock_ledger import clean_id, norm


class MasterError(Exception):
    """The master data can't be read or loaded (message is user-facing)."""


# Header text exactly as the Lark Base export writes it.
HEADERS = {
    "name": "Product Name",
    "sku": "Product ID",
    "alt_id": "Alt ID",
    "clients": "Clients",
    "unit": "Unit",
    "storage_unit": "Storage Unit",
    "per_carton": "Unit per Storage Unit",
    "weight": "น้ำหนักต่อชิ้น (kg)",
}
REQUIRED = ("name", "sku", "clients", "per_carton")
HEADER_SEARCH_ROWS = 5
MASTER_FILE = "stock_products.json"
PREVIOUS_FILE = "stock_products.previous.json"

# Short labels for the report. Anything not listed is shown as the master's own word, lower-cased.
UNIT_ABBREVIATIONS = {"carton": "ctn", "piece": "pcs", "bottle": "btl"}


# ─────────────────────────────────────────────
#  DATA MODEL
# ─────────────────────────────────────────────

@dataclass
class Product:
    sku: str
    name: str
    clients: list
    pieces_per_carton: Optional[int]
    storage_unit: str = "Carton"
    unit: str = "Piece"
    weight_per_piece: Optional[float] = None       # kg
    alt_id: str = ""
    row: int = 0                                   # row in the uploaded sheet, for messages


@dataclass
class ProductMaster:
    products: dict = field(default_factory=dict)   # {sku: Product}, in file order
    problems: list = field(default_factory=list)   # things to fix in Lark Base; the data is still usable
    updated_at: Optional[str] = None               # UTC, ISO
    source_filename: Optional[str] = None

    def lookup(self, sku: str) -> Optional[Product]:
        return self.products.get(clean_id(sku))

    def clients(self) -> list:
        """Every client that owns at least one product, sorted. This is the dropdown."""
        names = {c for p in self.products.values() for c in p.clients}
        return sorted(names, key=str.casefold)

    def zero_hint(self, sku: str) -> Optional[str]:
        """If `sku` isn't known but differs from a known ID only by leading zeros, say so.
        A hint only: we never match on it."""
        bare = sku.lstrip("0")
        if not bare or not sku.isdigit():
            return None
        for known in self.products:
            if known != sku and known.isdigit() and known.lstrip("0") == bare:
                return known
        return None

    def to_dict(self) -> dict:
        return {"updated_at": self.updated_at, "source_filename": self.source_filename,
                "problems": self.problems, "products": [asdict(p) for p in self.products.values()]}

    @classmethod
    def from_dict(cls, data: dict) -> "ProductMaster":
        try:
            products = {p["sku"]: Product(**p) for p in data["products"]}
            return cls(products=products, problems=list(data.get("problems", [])),
                       updated_at=data.get("updated_at"), source_filename=data.get("source_filename"))
        except (KeyError, TypeError) as e:
            raise MasterError(f"The stored product master data is damaged ({e!r}). Upload it again.")


# ─────────────────────────────────────────────
#  READING THE LARK BASE EXPORT
# ─────────────────────────────────────────────

def _split_clients(value) -> list:
    out: list = []
    for part in re.split(r"[,;\n]", str(value or "")):
        name = norm(part)
        if name and name not in out:
            out.append(name)
    return out


def _carton_size(value) -> tuple:
    """(pieces per carton or None, problem text or None) from text like '30 pieces per carton'."""
    text = norm(value)
    if not text:
        return None, "no carton size ('Unit per Storage Unit' is empty)"
    m = re.match(r"(\d+)\b", text)
    if not m:
        return None, f"carton size '{text}' can't be read (expected e.g. '30 pieces per carton')"
    n = int(m.group(1))
    if n <= 0:
        return None, f"carton size '{text}' must be more than 0"
    return n, None


def _find_header(wb):
    for ws in wb.worksheets:
        for r in range(1, HEADER_SEARCH_ROWS + 1):
            cells: dict = {}
            for c in range(1, ws.max_column + 1):
                text = norm(ws.cell(r, c).value)
                if text and text not in cells:
                    cells[text] = c
            if norm(HEADERS["sku"]) in cells and norm(HEADERS["name"]) in cells:
                col = {k: cells[norm(h)] for k, h in HEADERS.items() if norm(h) in cells}
                missing = [HEADERS[k] for k in REQUIRED if k not in col]
                if missing:
                    raise MasterError(f"Sheet '{ws.title}' is missing column(s): {', '.join(missing)}. "
                                      "Export the product list from Lark Base with those columns included.")
                return ws, r, col
    raise MasterError(
        "This doesn't look like the product master data: no header row with "
        f"'{HEADERS['sku']}' and '{HEADERS['name']}' was found. Export the 3PL product list from Lark Base.")


def parse_master_workbook(xlsx_path: str, source_filename: Optional[str] = None) -> ProductMaster:
    """Read the Lark Base product export. Raises MasterError if it isn't the product list;
    everything fixable in Lark Base is reported in master.problems and the rest still loads."""
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        raise MasterError(f"Could not open the workbook: {e}")
    try:
        ws, header_row, col = _find_header(wb)
        rows = list(ws.iter_rows(min_row=header_row + 1, values_only=True))
    finally:
        wb.close()

    def cell(values, key):
        i = col.get(key)
        return values[i - 1] if i and i - 1 < len(values) else None

    master = ProductMaster(updated_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
                           source_filename=source_filename)
    no_carton: list = []
    for offset, values in enumerate(rows):
        r = header_row + 1 + offset
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        name, sku = norm(cell(values, "name")), clean_id(cell(values, "sku"))
        alt = str(cell(values, "alt_id") or "").strip()
        where = f"Row {r} ({name or 'no name'})"

        if not sku:
            master.problems.append(f"{where}: no Product ID" + (f" (Alt ID: {norm(alt)})" if alt else "")
                                   + ". It can't be matched to a stock report until one is added.")
            continue
        per_carton, carton_problem = _carton_size(cell(values, "per_carton"))
        if carton_problem and per_carton is None and not norm(cell(values, "per_carton")):
            no_carton.append(sku)                   # one grouped message at the end; a whole client can lack them
        elif carton_problem:
            master.problems.append(f"{where}, SKU {sku}: {carton_problem}. Its quantities will show in pieces only.")
        clients = _split_clients(cell(values, "clients"))
        if not clients:
            master.problems.append(f"{where}, SKU {sku}: no client. It won't be offered to any client.")
        weight = None
        raw_weight = cell(values, "weight")
        if raw_weight not in (None, ""):
            try:
                weight = float(str(raw_weight).replace(",", ""))
            except ValueError:
                master.problems.append(f"{where}, SKU {sku}: weight per piece '{raw_weight}' is not a number.")

        product = Product(sku=sku, name=name, clients=clients, pieces_per_carton=per_carton,
                          storage_unit=norm(cell(values, "storage_unit")) or "Carton",
                          unit=norm(cell(values, "unit")) or "Piece",
                          weight_per_piece=weight, alt_id=alt, row=r)
        if sku in master.products:
            first = master.products[sku]
            note = (f"carton size {first.pieces_per_carton} vs {per_carton}"
                    if first.pieces_per_carton != per_carton else "same values")
            master.problems.append(f"{where}: Product ID {sku} is listed twice (first on row {first.row}; {note}). "
                                   "The first one is used.")
            continue
        master.products[sku] = product
    if no_carton:
        master.problems.append(f"{len(no_carton)} product(s) have no carton size ('Unit per Storage Unit' is empty), "
                               f"so their quantities will show in pieces only: {', '.join(no_carton)}.")
    return master


# ─────────────────────────────────────────────
#  CLIENT CHECK
# ─────────────────────────────────────────────

@dataclass
class ClientCheck:
    verdict: str                                   # "ok" | "warn" | "stop"
    messages: list = field(default_factory=list)
    matched: list = field(default_factory=list)    # SKUs that belong to the selected client
    unknown: list = field(default_factory=list)    # SKUs not in the master data
    other_clients: dict = field(default_factory=dict)   # {sku: [clients]} for SKUs that belong to someone else
    no_carton_size: list = field(default_factory=list)  # matched SKUs we can't convert to cartons

    @property
    def blocking(self) -> bool:
        return self.verdict == "stop"


def check_client(master: ProductMaster, client: str, skus: Iterable[str]) -> ClientCheck:
    """Does this file's stock belong to the client that was selected?

    The WMS file never names its client, so the selection is only a label. Filing one client's
    stock under another client's name is the mistake that matters most, so:
      stop  - none of the file's known SKUs belong to the selected client (or none are known at all)
      warn  - some SKUs belong to someone else, are unknown, or have no carton size
      ok    - everything matches
    """
    seen: list = []
    for s in skus:
        s = clean_id(s)
        if s and s not in seen:
            seen.append(s)

    if client not in master.clients():
        return ClientCheck("stop", [f"'{client}' isn't a client in the product master data."])

    result = ClientCheck("ok")
    for s in seen:
        p = master.lookup(s)
        if p is None:
            result.unknown.append(s)
        elif client in p.clients:
            result.matched.append(s)
            if not p.pieces_per_carton:
                result.no_carton_size.append(s)
        else:
            result.other_clients[s] = list(p.clients)

    known = len(result.matched) + len(result.other_clients)
    if known == 0:
        result.verdict = "stop"
        result.messages.append(
            "None of the SKUs in this file are in the product master data, so the client can't be checked. "
            "Update the master data, or check that this is the right file.")
    elif not result.matched:
        owners: dict = {}
        for owner_list in result.other_clients.values():
            for o in owner_list:
                owners[o] = owners.get(o, 0) + 1
        who = ", ".join(f"{o} ({n} SKU{'s' if n != 1 else ''})" for o, n in sorted(owners.items(), key=lambda kv: -kv[1]))
        result.verdict = "stop"
        result.messages.append(f"None of the SKUs in this file belong to {client}. They belong to {who}. "
                               "Check that you chose the right client.")
    elif result.other_clients:
        result.verdict = "warn"
        listed = ", ".join(f"{s} ({'/'.join(c)})" for s, c in result.other_clients.items())
        result.messages.append(f"{len(result.other_clients)} SKU(s) in this file belong to a different client: {listed}.")

    if result.unknown:
        if result.verdict == "ok":
            result.verdict = "warn"
        parts = []
        for s in result.unknown:
            hint = master.zero_hint(s)
            parts.append(f"{s} (did you mean {hint}? the file may have lost a leading zero)" if hint else s)
        if known:
            result.messages.append(f"{len(result.unknown)} SKU(s) are not in the product master data and will show "
                                   f"in pieces only: {', '.join(parts)}.")
        else:
            result.messages.append("Not found: " + ", ".join(parts) + ".")
    if result.no_carton_size:
        if result.verdict == "ok":
            result.verdict = "warn"
        result.messages.append(f"{len(result.no_carton_size)} SKU(s) have no carton size in the master data, so "
                               f"they show in pieces only: {', '.join(result.no_carton_size)}.")
    return result


# ─────────────────────────────────────────────
#  CARTONS + LEFTOVER PIECES
# ─────────────────────────────────────────────

@dataclass
class CartonQty:
    cartons: Optional[int]          # None when the carton size is unknown
    pieces: int                     # leftover pieces (all of them when cartons is None)
    total: int                      # total pieces
    per_carton: Optional[int]


def to_cartons(pieces: float, per_carton: Optional[int]) -> CartonQty:
    """Whole cartons plus leftover pieces: 1061 pieces at 30 per carton -> 35 cartons + 11 pieces.
    Pieces are rounded to the nearest whole piece (half rounds up); negatives keep their sign on both parts."""
    total = int(math.floor(abs(pieces) + 0.5)) * (1 if pieces >= 0 else -1)
    if not per_carton or per_carton <= 0:
        return CartonQty(None, total, total, None)
    sign = -1 if total < 0 else 1
    cartons, leftover = divmod(abs(total), per_carton)
    return CartonQty(sign * cartons if cartons else 0, sign * leftover if leftover else 0, total, per_carton)


def unit_labels(product: Optional[Product]) -> tuple:
    """(storage label, piece label) for the report, e.g. ('ctn', 'pcs') or ('crate', 'btl')."""
    def short(word: str, default: str) -> str:
        word = (word or default).strip().lower()
        return UNIT_ABBREVIATIONS.get(word, word)
    return (short(product.storage_unit if product else "", "carton"), short(product.unit if product else "", "piece"))


def format_cartons(q: CartonQty, storage_label: str = "ctn", unit_label: str = "pcs") -> str:
    """'35 ctn + 11 pcs'. A zero part is left out ('12 ctn', '9 pcs'); nothing at all is '0'.
    Unknown carton size: pieces only ('178 pcs'). Negative: '-2 ctn + 5 pcs'."""
    if q.cartons is None:
        return f"{q.total:,} {unit_label}"
    if q.total == 0:
        return "0"
    sign = "-" if q.total < 0 else ""
    parts = []
    if q.cartons:
        parts.append(f"{abs(q.cartons):,} {storage_label}")
    if q.pieces:
        parts.append(f"{abs(q.pieces):,} {unit_label}")
    return sign + " + ".join(parts)


def carton_view(summary, product: Optional[Product]) -> dict:
    """The four quantities of a stock_ledger.SkuSummary as CartonQty:
    {"opening", "in", "out", "closing"}. A SKU with no product (or no carton size) stays in pieces."""
    per = product.pieces_per_carton if product else None
    return {"opening": to_cartons(summary.opening_qty, per), "in": to_cartons(summary.in_qty, per),
            "out": to_cartons(summary.out_qty, per), "closing": to_cartons(summary.closing_qty, per)}


# ─────────────────────────────────────────────
#  STORAGE
# ─────────────────────────────────────────────

class MasterStore:
    """The master data on disk, in `directory` (on the volume: DATA_DIR/masterdata).
    Not under uploads/ or output/, so the cleanup sweep never touches it.
    Saving keeps the previous copy; there is no login, so that copy is the safety net."""

    def __init__(self, directory):
        self.directory = Path(directory)
        self.path = self.directory / MASTER_FILE
        self.previous_path = self.directory / PREVIOUS_FILE

    def load(self) -> Optional[ProductMaster]:
        """The stored master, or None if none has been uploaded yet."""
        return self._read(self.path)

    def load_previous(self) -> Optional[ProductMaster]:
        return self._read(self.previous_path)

    @staticmethod
    def _read(path: Path) -> Optional[ProductMaster]:
        if not path.exists():
            return None
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as e:
            raise MasterError(f"The stored product master data can't be read ({e}). Upload it again.")
        return ProductMaster.from_dict(data)

    def save(self, master: ProductMaster) -> None:
        self.directory.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            shutil.copy2(self.path, self.previous_path)
        self._write(self.path, master)

    def restore_previous(self) -> bool:
        """Put the previous copy back. Returns False if there isn't one."""
        previous = self.load_previous()
        if previous is None:
            return False
        self._write(self.path, previous)
        return True

    def _write(self, path: Path, master: ProductMaster) -> None:
        """Write to a temp file then rename, so a crash can never leave half a file."""
        fd, tmp = tempfile.mkstemp(dir=self.directory, suffix=".tmp")
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                json.dump(master.to_dict(), f, ensure_ascii=False, indent=1)
            os.replace(tmp, path)
        except BaseException:
            Path(tmp).unlink(missing_ok=True)
            raise


def diff_masters(old: Optional[ProductMaster], new: ProductMaster) -> dict:
    """What an upload would change, so a person sees it before relying on it."""
    old_products = old.products if old else {}
    added = [s for s in new.products if s not in old_products]
    removed = [s for s in old_products if s not in new.products]
    carton_changed = [(s, old_products[s].pieces_per_carton, new.products[s].pieces_per_carton)
                      for s in new.products if s in old_products
                      and old_products[s].pieces_per_carton != new.products[s].pieces_per_carton]
    client_changed = [(s, old_products[s].clients, new.products[s].clients)
                      for s in new.products if s in old_products and old_products[s].clients != new.products[s].clients]
    return {"added": added, "removed": removed, "carton_changed": carton_changed, "client_changed": client_changed}
